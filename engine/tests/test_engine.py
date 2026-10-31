"""湛卢引擎单元 + API 冒烟测试。运行：engine/.venv/Scripts/python -m unittest discover -s tests"""

import io
import json
import os
import tempfile
import threading
import unittest
import urllib.request
import zipfile
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import _zlenv  # noqa: F401 — 固定测试数据目录（须先于 zhanlu 导入）

TMP = _zlenv.__dict__.get("_TMP") or Path(os.environ["ZHANLU_HOME"]).parent

from pypdf import PdfReader, PdfWriter  # noqa: E402

from zhanlu.core import ai, ofd, pdfops, tables, topdf, export  # noqa: E402
from zhanlu import secure  # noqa: E402


def make_pdf(path: Path, pages: int = 6, outline: bool = False, text: bool = False) -> Path:
    writer = PdfWriter()
    for i in range(pages):
        page = writer.add_blank_page(width=595, height=842)
        if text:
            from reportlab.pdfgen import canvas
            buffer = io.BytesIO()
            drawing = canvas.Canvas(buffer, pagesize=(595, 842))
            drawing.drawString(72, 750, f"Page {i + 1} sample text")
            drawing.save()
            buffer.seek(0)
            page.merge_page(PdfReader(buffer).pages[0])
    if outline:
        c1 = writer.add_outline_item("第一章 审计意见", 0)
        writer.add_outline_item("1.1 说明", 1, parent=c1)
        writer.add_outline_item("第二章 财务报表", 2)
    with open(path, "wb") as f:
        writer.write(f)
    return path


def make_table_pdf(path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    from zhanlu.core.topdf import cjk_font

    rows = [["项目", "本期金额", "上期金额"], ["营业收入", "531,208,933.87", "482,130,555.20"],
            ["净利润", "71,488,207.90", "56,511,067.32"]]
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    t = Table(rows)
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("FONTNAME", (0, 0), (-1, -1), cjk_font()),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
    ]))
    doc.build([t])
    return path


def make_ofd(path: Path) -> Path:
    ofd_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<OFD version="1.0"><DocBody>'
        '<DocInfo><DocID>TEST1</DocID><Title>测试 OFD 文档</Title><Author>湛卢</Author></DocInfo>'
        '<DocRoot>Doc_0/Document.xml</DocRoot>'
        "</DocBody></OFD>"
    )
    doc_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Document><Pages>'
        '<Page ID="P1" BaseLoc="Pages/P_0/P_0.xml"/>'
        '<Page ID="P2" BaseLoc="Pages/P_1/P_1.xml"/>'
        "</Pages></Document>"
    )
    def page_xml(text: str) -> str:
        return (
            '<?xml version="1.0" encoding="UTF-8"?>'
            f'<Page><Content><TextObject Boundary="20 20 170 20" Size="5"><TextCode X="0" Y="5">{text}</TextCode></TextObject></Content></Page>'
        )
    bm_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Bookmarks><Bookmark Title="第一章"><Point PageRef="P1"/></Bookmark>'
        '<Bookmark Title="第二章"><Point PageRef="P2"/></Bookmark></Bookmarks>'
    )
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("OFD.xml", ofd_xml)
        z.writestr("Doc_0/Document.xml", doc_xml)
        z.writestr("Doc_0/Pages/P_0/P_0.xml", page_xml("湛卢测试第一页文本"))
        z.writestr("Doc_0/Pages/P_1/P_1.xml", page_xml("第二页内容 Zhanlu"))
        z.writestr("Doc_0/Bookmark.xml", bm_xml)
    return path


class TestPdfOps(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = TMP / "pdfops"
        cls.dir.mkdir(parents=True, exist_ok=True)
        cls.src = make_pdf(cls.dir / "src.pdf", 6, outline=True)

    def test_parse_pages(self):
        self.assertEqual(pdfops.parse_pages("1-3,5", 6), [0, 1, 2, 4])
        self.assertEqual(pdfops.parse_pages("8-", 8), [7])
        self.assertEqual(pdfops.parse_pages(None, 6), None)
        with self.assertRaises(pdfops.OpError):
            pdfops.parse_pages("9", 6)

    def test_info_and_outline(self):
        info = pdfops.info(self.src)
        self.assertEqual(info["pages"], 6)
        self.assertFalse(info["encrypted"])
        nodes = pdfops.outline(self.src)
        self.assertEqual(nodes[0]["title"], "第一章 审计意见")
        self.assertEqual(nodes[0]["page"], 1)
        self.assertEqual(nodes[0]["children"][0]["page"], 2)
        self.assertEqual(nodes[1]["title"], "第二章 财务报表")
        md = pdfops.toc_markdown(nodes)
        self.assertIn("第一章 审计意见 · 1", md)
        self.assertIn("  - 1.1 说明 · 2", md)

    def test_merge(self):
        out = self.dir / "merged.pdf"
        pdfops.merge([{"path": self.src}, {"path": self.src, "pages": "2"}], out)
        self.assertEqual(len(PdfReader(str(out)).pages), 7)

    def test_split_ranges_every(self):
        outs = pdfops.split(self.src, self.dir / "s1", "ranges", ranges="1-2,3-4,5-6")
        self.assertEqual([len(PdfReader(str(o)).pages) for o in outs], [2, 2, 2])
        outs = pdfops.split(self.src, self.dir / "s2", "every", every=4)
        self.assertEqual([len(PdfReader(str(o)).pages) for o in outs], [4, 2])

    def test_delete_extract_rotate_move(self):
        out = self.dir / "del.pdf"
        pdfops.delete_pages(self.src, "2,4", out)
        self.assertEqual(len(PdfReader(str(out)).pages), 4)
        out = self.dir / "ext.pdf"
        pdfops.extract_pages(self.src, "1,3", out)
        self.assertEqual(len(PdfReader(str(out)).pages), 2)
        out = self.dir / "rot.pdf"
        pdfops.rotate_pages(self.src, "1-2", 90, out)
        self.assertEqual(PdfReader(str(out)).pages[0].rotation, 90)
        out = self.dir / "mv.pdf"
        pdfops.move_pages(self.src, "1", 3, out)
        self.assertEqual(len(PdfReader(str(out)).pages), 6)

    def test_insert_blank_and_pdf(self):
        out = self.dir / "blank.pdf"
        pdfops.insert_blank(self.src, 2, 2, out)
        self.assertEqual(len(PdfReader(str(out)).pages), 8)
        out = self.dir / "ins.pdf"
        pdfops.insert_pdf(self.src, self.src, "1", 3, out)
        self.assertEqual(len(PdfReader(str(out)).pages), 7)

    def test_encrypt_decrypt_roundtrip(self):
        enc = self.dir / "enc.pdf"
        pdfops.encrypt(self.src, enc, "mima123")
        reader = PdfReader(str(enc))
        self.assertTrue(reader.is_encrypted)
        dec = self.dir / "dec.pdf"
        pdfops.decrypt(enc, dec, "mima123")
        self.assertEqual(len(PdfReader(str(dec)).pages), 6)

    def test_render(self):
        outs = pdfops.render_pages(self.src, self.dir / "img", pages="1-2", dpi=72)
        self.assertEqual(len(outs), 2)
        self.assertTrue(outs[0].stat().st_size > 100)


class TestExtractAndGen(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = TMP / "gen"
        cls.dir.mkdir(parents=True, exist_ok=True)

    def test_tables(self):
        pdf = make_table_pdf(self.dir / "table.pdf")
        data = tables.extract(pdf)
        self.assertTrue(data, "应检测到表格")
        rows = data[0]["tables"][0]
        self.assertEqual(rows[0][0], "项目")
        self.assertEqual(rows[2][1], "71,488,207.90")
        xlsx = tables.export_xlsx(data, self.dir / "t.xlsx")
        csvp = tables.export_csv(data, self.dir / "t.csv")
        self.assertTrue(xlsx.exists() and csvp.exists())
        self.assertIn("71,488,207.90", csvp.read_text(encoding="utf-8-sig"))

    def test_txt_md_pdf(self):
        out = self.dir / "txt.pdf"
        topdf.txt_to_pdf("湛卢测试段落一。\n\n第二段 Zhanlu test.", out, title="湛卢")
        reader = PdfReader(str(out))
        self.assertGreaterEqual(len(reader.pages), 1)
        text = reader.pages[0].extract_text() or ""
        self.assertIn("湛卢", text)
        out2 = self.dir / "md.pdf"
        topdf.md_to_pdf("# 标题\n\n正文 **加粗** 项\n\n- 列表一\n- 列表二\n\n```code```", out2, title="MD")
        self.assertTrue(out2.exists())

    def test_images_to_pdf(self):
        from PIL import Image

        imgs = []
        for color in ((200, 30, 30), (30, 30, 200)):
            p = self.dir / f"img_{color[0]}.png"
            Image.new("RGB", (60, 80), color).save(p)
            imgs.append(p)
        out = topdf.images_to_pdf(imgs, self.dir / "pics.pdf")
        self.assertEqual(len(PdfReader(str(out)).pages), 2)

    def test_pdf_to_docx_pptx(self):
        src = make_pdf(self.dir / "six.pdf", 3)
        out = export.pdf_to_docx(src, self.dir / "out.docx")
        self.assertTrue(out.stat().st_size > 1000)
        imgs = pdfops.render_pages(src, self.dir / "pp", dpi=72)
        out2 = export.pdf_to_pptx(imgs, self.dir / "out.pptx")
        self.assertTrue(out2.stat().st_size > 10000)


class TestConvertOfficeTargets(unittest.TestCase):
    """convert_pdf 新目标（docx/pptx/xlsx/csv）：预览说明、执行产物与报错。"""

    @classmethod
    def setUpClass(cls):
        cls.dir = TMP / "convert-office"
        cls.dir.mkdir(parents=True, exist_ok=True)

    def _ctx(self):
        from zhanlu.tasks import TaskContext

        return TaskContext({"id": "office-test", "outputs": [], "pending_outputs": [], "cancel": threading.Event()})

    def _run(self, source, params, pages_total):
        from zhanlu import jobs
        from zhanlu.tasks import TaskContext

        entry = {"id": "office-src", "path": str(source), "name": source.name}
        task = {"id": "office-test", "outputs": [], "pending_outputs": [], "cancel": threading.Event()}
        with patch.object(jobs, "_entry_path", return_value=(source, entry)), \
             patch.object(jobs, "_pages_of", return_value=pages_total):
            return jobs.run_convert_pdf({"id": entry["id"], "dir": str(self.dir), **params},
                                        TaskContext(task))

    def _preview(self, source, params, pages_total):
        from zhanlu import jobs

        entry = {"id": "office-src", "path": str(source), "name": source.name}
        with patch.object(jobs, "_entry_path", return_value=(source, entry)), \
             patch.object(jobs, "_pages_of", return_value=pages_total):
            return jobs.preview_convert_pdf({"id": entry["id"], "dir": str(self.dir), **params})

    def test_docx_export_contains_selected_page_text(self):
        source = make_pdf(self.dir / "doc.pdf", 3, text=True)
        result = self._run(source, {"target": "docx", "prefix": "doc", "pages": "2"}, 3)
        out = Path(result["files"][0]["path"])
        self.assertEqual(out.suffix, ".docx")
        from docx import Document

        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        self.assertIn("Page 2 sample text", text)
        self.assertNotIn("Page 1 sample text", text)

    def test_docx_export_sanitizes_control_characters(self):
        """提取文本混入 XML 非法控制字符时照常导出，不报 ValueError（真实 PDF 的 ToUnicode
        映射会让 \x0b 等直达文本层；Helvetica/WinAnsi fixture 会把控制字节渲染成 ■，无法复现）。"""
        from zhanlu.core import export as export_mod

        src = make_pdf(self.dir / "plain.pdf", 1)
        out = self.dir / "ctrl.docx"
        with patch.object(export_mod, "extract_text", return_value=["bad\x0bchars\x01here\x02tail"]):
            export.pdf_to_docx(src, out)
        from docx import Document

        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        self.assertIn("charsheretail", text)  # \x0b 按 splitlines 成行属预期，\x01/\x02 被清理
        self.assertNotRegex(text, r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

    def test_pptx_export_slide_per_page_and_dpi_guard(self):
        source = make_pdf(self.dir / "ppt.pdf", 2, text=True)
        result = self._run(source, {"target": "pptx", "prefix": "ppt", "dpi": 96}, 2)
        out = Path(result["files"][0]["path"])
        self.assertEqual(out.suffix, ".pptx")
        from pptx import Presentation

        self.assertEqual(len(Presentation(str(out)).slides), 2)
        for slide in Presentation(str(out)).slides:
            picture = slide.shapes[0]
            self.assertAlmostEqual(picture.width / picture.height, 595 / 842, delta=0.002)
            self.assertGreater(picture.left, 0)
            self.assertGreaterEqual(picture.top, 0)
        with self.assertRaisesRegex(pdfops.OpError, "DPI"):
            self._run(source, {"target": "pptx", "prefix": "ppt", "dpi": 120}, 2)

    def test_xlsx_export_and_empty_preview_block(self):
        source = make_table_pdf(self.dir / "table.pdf")
        prev = self._preview(source, {"target": "xlsx", "prefix": "tbl"}, 1)
        self.assertTrue(prev["ok"])
        self.assertIn("检测到", prev["summary"])
        result = self._run(source, {"target": "xlsx", "prefix": "tbl"}, 1)
        out = Path(result["files"][0]["path"])
        import openpyxl

        wb = openpyxl.load_workbook(str(out))
        self.assertEqual(wb.sheetnames[0], "P1T1")
        self.assertEqual(wb[wb.sheetnames[0]].cell(1, 1).value, "项目")

        blank = make_pdf(self.dir / "blank.pdf", 1)
        prev_empty = self._preview(blank, {"target": "xlsx", "prefix": "tbl-blank"}, 1)
        self.assertFalse(prev_empty["ok"])
        self.assertIn("未检测到表格", prev_empty["error"])
        with self.assertRaisesRegex(pdfops.OpError, "未检测到表格"):
            self._run(blank, {"target": "xlsx", "prefix": "tbl-blank"}, 1)
        self.assertFalse((self.dir / "tbl-blank.xlsx").exists())

    def test_csv_export(self):
        source = make_table_pdf(self.dir / "table2.pdf")
        result = self._run(source, {"target": "csv", "prefix": "tbl2"}, 1)
        out = Path(result["files"][0]["path"])
        self.assertEqual(out.suffix, ".csv")
        self.assertIn("71,488,207.90", out.read_text(encoding="utf-8-sig"))

    def test_table_after_preview_sample_can_be_exported(self):
        """前八页没有表格不能阻止导出第九页的真实表格。"""
        source = self.dir / "late-table.pdf"
        writer = PdfWriter()
        writer.append(make_pdf(self.dir / "intro.pdf", 8, text=True))
        writer.append(make_table_pdf(self.dir / "appendix.pdf"))
        with source.open("wb") as stream:
            writer.write(stream)
        for target in ("xlsx", "csv"):
            prev = self._preview(source, {"target": target, "prefix": "late"}, 9)
            self.assertTrue(prev["ok"])
            self.assertIn("检测到 1 张表格", prev["summary"])
            result = self._run(source, {"target": target, "prefix": "late"}, 9)
            self.assertTrue(Path(result["files"][0]["path"]).exists())

    def test_unknown_target_message(self):
        source = make_pdf(self.dir / "bad.pdf", 1)
        with self.assertRaisesRegex(pdfops.OpError, "Markdown"):
            self._run(source, {"target": "exe", "prefix": "bad"}, 1)


class TestOfd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = TMP / "ofd"
        cls.dir.mkdir(parents=True, exist_ok=True)
        cls.src = make_ofd(cls.dir / "test.ofd")

    def test_info_outline_text(self):
        info = ofd.info(self.src)
        self.assertEqual(info["pages"], 2)
        self.assertEqual(info["doc"]["Title"], "测试 OFD 文档")
        nodes = info["outline"]
        self.assertEqual([n["title"] for n in nodes], ["第一章", "第二章"])
        self.assertEqual(nodes[1]["page"], 2)
        pages = ofd.text(self.src)
        self.assertIn("湛卢测试第一页文本", pages[0])
        self.assertIn("Zhanlu", pages[1])

    def test_ofd_to_pdf(self):
        out = self.dir / "conv.pdf"
        result = ofd.to_pdf(self.src, out)
        reader = PdfReader(str(out))
        self.assertEqual(len(reader.pages), 2)
        self.assertIn("版式级", result["note"])
        self.assertIn("湛卢测试第一页文本", reader.pages[0].extract_text())

    def test_layout_reports_failed_objects(self):
        with patch.object(ofd, "_draw_text", side_effect=RuntimeError("绘制失败")):
            result = ofd.to_pdf(self.src, self.dir / "partial.pdf")
        self.assertIn("部分文字未渲染", result["note"])


class TestOfdTasks(unittest.TestCase):
    """OFD 预览缓存与转 PDF 任务：交付、缓存命中与文字重建回落。"""

    @classmethod
    def setUpClass(cls):
        cls.dir = TMP / "ofd-tasks"
        cls.dir.mkdir(parents=True, exist_ok=True)

    def _ctx(self):
        from zhanlu.tasks import TaskContext

        return TaskContext({"id": "ofd-task", "outputs": [], "pending_outputs": [], "cancel": threading.Event()})

    def _entry(self, source):
        return {"id": "ofd-src", "path": str(source), "name": source.name, "kind": "ofd"}

    def test_preview_convert_delivers_pdf_with_cache(self):
        from zhanlu import jobs
        from zhanlu.tasks import TaskContext

        source = make_ofd(self.dir / "doc.ofd")
        entry = self._entry(source)
        with patch.object(jobs, "_entry_path", return_value=(source, entry)):
            prev = jobs.preview_ofd_convert({"id": "ofd-src", "dir": str(self.dir)})
            self.assertTrue(prev["ok"])
            self.assertEqual(prev["outputs"][0]["name"], "doc-转换.pdf")
            ctx = self._ctx()
            result = jobs.run_ofd_convert({"id": "ofd-src", "dir": str(self.dir)}, ctx)
            out = Path(result["files"][0]["path"])
            self.assertEqual(out.suffix, ".pdf")
            self.assertEqual(len(PdfReader(str(out)).pages), 2)
            # 预览缓存命中（同内容第二次转换走缓存）
            with patch.object(jobs, "_check_run_target"), patch.object(jobs, "_deliver", return_value=entry), \
                 patch.object(jobs.store, "set_pages"):
                second = jobs.run_ofd_convert({"id": "ofd-src", "name": "doc2", "dir": str(self.dir)},
                                              TaskContext({"id": "ofd-task2", "outputs": [], "pending_outputs": [], "cancel": threading.Event()}))
            self.assertEqual(result["note"], second["note"])

    def test_ofd_preview_registers_internal_file(self):
        from zhanlu import jobs

        source = make_ofd(self.dir / "doc2.ofd")
        entry = self._entry(source)
        with patch.object(jobs, "_entry_path", return_value=(source, entry)):
            result = jobs.run_ofd_preview({"id": "ofd-src"}, self._ctx())
        cached = result["file"]
        self.assertEqual(cached["kind"], "pdf")
        self.assertFalse(cached.get("recent", True), "预览缓存不应进入最近文件")
        self.assertEqual(cached["pages"], 2)
        self.assertTrue(Path(cached["path"]).exists())

    def test_layout_failure_falls_back_to_text_rebuild(self):
        from zhanlu import jobs

        source = make_ofd(self.dir / "doc3.ofd")
        entry = self._entry(source)
        with patch.object(jobs, "_entry_path", return_value=(source, entry)), \
             patch.object(jobs.ofd, "to_pdf", side_effect=RuntimeError("模拟版式失败")):
            result = jobs.run_ofd_preview({"id": "ofd-src"}, self._ctx())
        self.assertIn("文字重建版", result["note"])
        self.assertTrue(Path(result["file"]["path"]).exists())
        with patch.object(jobs.ofd, "to_pdf") as convert:
            _, cached_note = jobs._ofd_cached_pdf(source, ctx=self._ctx())
        convert.assert_not_called()
        self.assertIn("文字重建版", cached_note)

    def test_non_ofd_rejected(self):
        from zhanlu import jobs

        source = make_pdf(self.dir / "plain.pdf", 1)
        entry = {"id": "pdf-src", "path": str(source), "name": source.name, "kind": "pdf"}
        with patch.object(jobs, "_entry_path", return_value=(source, entry)):
            with self.assertRaisesRegex(pdfops.OpError, "仅支持 OFD"):
                jobs.preview_ofd_convert({"id": "pdf-src", "dir": str(self.dir)})

    def test_concurrent_ofd_cache_builds_once(self):
        from zhanlu import jobs
        import time

        source = make_ofd(self.dir / "concurrent.ofd")
        barrier = threading.Barrier(2)
        def convert(_source, out, **kwargs):
            time.sleep(0.05)
            make_pdf(out, 1, text=True)
            return {"note": "文字重建版（测试）"}
        def load():
            barrier.wait()
            return jobs._ofd_cached_pdf(source, ctx=self._ctx())
        with tempfile.TemporaryDirectory(dir=self.dir) as cache_dir, \
             patch.object(jobs.config, "CACHE_OFD_DIR", Path(cache_dir)), \
             patch.object(jobs.ofd, "to_pdf", side_effect=convert) as build:
            with ThreadPoolExecutor(max_workers=2) as pool:
                results = list(pool.map(lambda _: load(), range(2)))
            self.assertEqual(build.call_count, 1)
            self.assertEqual(results[0], results[1])
            self.assertTrue(results[0][0].exists())


class TestAiStore(unittest.TestCase):
    def test_obfuscate_roundtrip(self):
        secret = "sk-test-1234567890"
        token = ai.obfuscate(secret)
        self.assertNotEqual(token, secret)
        self.assertEqual(ai.deobfuscate(token), secret)

    def test_windows_key_protection_fails_closed(self):
        with patch.object(secure, "_is_windows", True), patch.object(
            secure, "_dpapi_protect", side_effect=OSError("DPAPI unavailable")
        ):
            with self.assertRaisesRegex(OSError, "DPAPI unavailable"):
                secure.protect_secret("sk-test-failure")

    @unittest.skipUnless(os.name == "nt", "Windows DPAPI only")
    def test_windows_dpapi_roundtrip(self):
        secret = "sk-test-dpapi-roundtrip"
        encrypted = secure.protect_secret(secret)
        self.assertTrue(__import__("base64").b64decode(encrypted).startswith(b"DP1|"))
        self.assertEqual(secure.unprotect_secret(encrypted), secret)

    def test_save_list_configs(self):
        payload = {
            "configs": [{"id": "c1", "name": "测试配置", "protocol": "openai",
                         "baseurl": "https://api.example.com/v1", "model": "test-model",
                         "key": "sk-abc", "temperature": 0.7, "timeout": 60, "headers": {}}],
            "active": "c1",
        }
        result = ai.save_configs(payload)
        self.assertTrue(result["ok"])
        listed = ai.list_configs()
        self.assertEqual(len(listed), 1)
        self.assertFalse("key" in listed[0])
        self.assertTrue(listed[0]["has_key"])
        full = ai.active_config_full()
        self.assertEqual(full["key_plain"], "sk-abc")

    def test_unreadable_saved_key_is_not_reported_as_configured(self):
        saved = {"ai": {"configs": [{"id": "broken", "key": "unreadable-ciphertext"}]}}
        with patch.object(ai, "_settings", return_value=saved), patch.object(ai, "deobfuscate", return_value=""):
            listed = ai.list_configs()
        self.assertFalse(listed[0]["has_key"])
        self.assertNotIn("key", listed[0])


class TestStoreConcurrency(unittest.TestCase):
    def test_upload_dedupe_does_not_inherit_native_source_directory(self):
        from zhanlu import jobs
        from zhanlu.store import Store, source_dir_of

        store = Store()
        native = make_pdf(TMP / "native-identical.pdf", 1)
        native_entry = store.add_path(native)
        first_tmp = TMP / "upload-identical-1.tmp"
        first_tmp.write_bytes(native.read_bytes())
        uploaded = store.add_bytes_from(first_tmp, "upload-identical.pdf", dedupe=True)
        self.assertNotEqual(uploaded["id"], native_entry["id"])
        self.assertIsNone(source_dir_of(uploaded))
        with self.assertRaisesRegex(pdfops.OpError, "可信的本机源路径"):
            jobs._resolve_dir(jobs.SOURCE_DIR_OPT, [uploaded])
        second_tmp = TMP / "upload-identical-2.tmp"
        second_tmp.write_bytes(native.read_bytes())
        duplicate = store.add_bytes_from(second_tmp, "upload-again.pdf", dedupe=True)
        self.assertEqual(duplicate["id"], uploaded["id"])
        self.assertIsNone(source_dir_of(duplicate))
        self.assertFalse(second_tmp.exists())

    def test_parallel_identical_uploads_share_one_inbox_entry(self):
        from zhanlu.store import Store, source_dir_of

        store = Store()
        paths = [TMP / f"upload-parallel-{i}.tmp" for i in range(8)]
        for path in paths:
            path.write_bytes(b"unique parallel upload")
        with ThreadPoolExecutor(max_workers=8) as pool:
            entries = list(pool.map(lambda path: store.add_bytes_from(path, "parallel.pdf", dedupe=True), paths))
        self.assertEqual(len({entry["id"] for entry in entries}), 1)
        self.assertIsNone(source_dir_of(entries[0]))

    def test_parallel_output_previews_do_not_share_probe_file(self):
        from zhanlu.jobs import _dir_writable

        target = TMP / "parallel-preview-output"
        with ThreadPoolExecutor(max_workers=8) as pool:
            results = list(pool.map(lambda _: _dir_writable(target), range(24)))
        self.assertEqual(results, [None] * 24)
        self.assertFalse(list(target.glob(".zl-write-probe-*")))

    @unittest.skipUnless(os.name == "nt", "Windows file replacement race")
    def test_atomic_json_retries_transient_access_denied(self):
        from zhanlu import config

        dest = TMP / "retry-registry.json"
        original = os.replace
        attempts = 0

        def transient_replace(src, dst):
            nonlocal attempts
            attempts += 1
            if attempts == 1:
                raise PermissionError(5, "file temporarily held")
            return original(src, dst)

        with patch("zhanlu.config.os.replace", side_effect=transient_replace):
            config.save_json(dest, {"ok": True})
        self.assertEqual(attempts, 2)
        self.assertEqual(config.load_json(dest, {}), {"ok": True})

    def test_concurrent_registration_keeps_all_entries(self):
        from zhanlu import config
        from zhanlu.store import Store

        store = Store()
        with ThreadPoolExecutor(max_workers=8) as pool:
            entries = list(pool.map(lambda i: store.add_bytes(b"sample", f"parallel-{i}.txt"), range(20)))
        ids = {entry["id"] for entry in entries}
        self.assertTrue(ids.issubset({entry["id"] for entry in store.list()}))
        self.assertTrue(ids.issubset(config.load_json(config.REGISTRY_FILE, {}).keys()))


class TestTaskCancellation(unittest.TestCase):
    def test_protected_pdf_page_organization_is_blocked_in_preview_and_run(self):
        from pypdf.constants import UserAccessPermissions
        from zhanlu import jobs
        from zhanlu.tasks import TaskContext

        source = TMP / "protected-organize.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.encrypt("", owner_password="owner", permissions_flag=UserAccessPermissions.PRINT)
        with source.open("wb") as stream:
            writer.write(stream)
        entry = {"id": "protected", "path": str(source), "name": source.name}
        params = {"id": "protected", "plan": [{"page": 1, "rotate": 90}], "dir": str(TMP)}
        ctx = TaskContext({"id": "protected", "outputs": [], "pending_outputs": [], "cancel": threading.Event()})
        with patch.object(jobs, "_entry_path", return_value=(source, entry)), \
             patch.object(jobs, "_pages_of", return_value=1):
            with self.assertRaisesRegex(pdfops.OpError, "限制修改"):
                jobs.preview_organize(params)
            with self.assertRaisesRegex(pdfops.OpError, "限制修改"):
                jobs.run_organize(params, ctx)

    def test_delivery_conflict_never_overwrites_and_numbered_retry_succeeds(self):
        from zhanlu import jobs
        from zhanlu.tasks import TaskContext

        directory = TMP / "delivery-race"
        directory.mkdir(exist_ok=True)
        source = directory / "source.txt"
        source.write_text("task output", encoding="utf-8")
        target = directory / "output.txt"
        task = {"outputs": [], "pending_outputs": [], "cancel": threading.Event()}
        ctx = TaskContext(task)
        publish_name = "rename" if os.name == "nt" else "link"
        original_publish = getattr(os, publish_name)
        def compete(staged, final):
            if Path(final) == target:
                target.write_text("other user file", encoding="utf-8")
            return original_publish(staged, final)
        with patch(f"zhanlu.jobs.os.{publish_name}", side_effect=compete):
            result = jobs._deliver_numbered(source, directory, "output.txt", ctx)
        self.assertEqual(target.read_text(encoding="utf-8"), "other user file")
        self.assertEqual(Path(result["path"]).name, "output（1）.txt")
        self.assertEqual(Path(result["path"]).read_text(encoding="utf-8"), "task output")
        self.assertEqual(len(task["pending_outputs"]), 1)

    def test_interrupted_reservation_does_not_remove_same_content_foreign_file(self):
        from zhanlu import jobs
        from zhanlu.tasks import TaskContext

        directory = TMP / "delivery-identity"
        directory.mkdir(exist_ok=True)
        staged = directory / "staged.txt"
        target = directory / "target.txt"
        staged.write_text("same content", encoding="utf-8")
        task = {"outputs": [], "pending_outputs": [], "cancel": threading.Event()}
        TaskContext(task).reserve_output(staged, target, staged=True)
        target.write_text("same content", encoding="utf-8")
        self.assertEqual(jobs._cleanup_pending(task["pending_outputs"]), task["pending_outputs"])
        self.assertTrue(target.exists())
        self.assertFalse(staged.exists())

    def test_txt_conversion_rejects_empty_and_keeps_text_pages(self):
        from zhanlu import jobs
        from zhanlu.tasks import TaskContext

        source = make_pdf(TMP / "txt-selected.pdf", 2, text=True)
        entry = {"id": "txt-selected", "path": str(source), "name": source.name}
        task = {"id": "txt-selected", "outputs": [], "pending_outputs": [], "cancel": threading.Event()}
        ctx = TaskContext(task)
        params = {"id": entry["id"], "target": "txt", "prefix": "txt-selected", "dir": str(TMP)}
        with patch.object(jobs, "_entry_path", return_value=(source, entry)), \
             patch.object(jobs, "_pages_of", return_value=2):
            result = jobs.run_convert_pdf(params, ctx)
            self.assertIn("Page 1 sample text", Path(result["files"][0]["path"]).read_text(encoding="utf-8"))
            blank = make_pdf(TMP / "txt-blank.pdf", 2)
            with patch.object(jobs, "_entry_path", return_value=(blank, entry)):
                with self.assertRaisesRegex(pdfops.OpError, "OCR"):
                    jobs.run_convert_pdf({**params, "prefix": "txt-blank"}, ctx)
            self.assertFalse((TMP / "txt-blank.txt").exists())

    def test_ocr_page_selection_uses_visible_numbers_without_truncation(self):
        from zhanlu import jobs

        class Context:
            _task = {"cancel": threading.Event()}
            temp_root = TMP

            def check(self):
                pass

            def progress(self, *_args):
                pass

            def output(self, _entry):
                pass

        entry = {"id": "ocr-test", "path": str(TMP / "ocr-test.pdf"), "name": "ocr-test.pdf"}
        with patch.object(jobs, "_entry_path", return_value=(Path(entry["path"]), entry)), \
             patch.object(jobs, "_pages_of", return_value=5) as pages_of, \
             patch.object(jobs, "_check_run_target"), \
             patch.object(jobs, "_deliver", return_value=entry), \
             patch.object(jobs.store, "set_pages"), \
             patch.object(jobs.ocrpdf, "ocr_pdf", return_value={"checked": 0}) as ocr, \
             patch.object(jobs.pdfops, "text_page_stats", return_value={"text_pages": []}):
            jobs.run_ocr({"id": entry["id"], "pages": "1-3", "dir": str(TMP)}, Context())
            self.assertEqual(ocr.call_args.kwargs["pages"], "1,2,3")

            pages_of.return_value = 4001
            jobs.run_ocr({"id": entry["id"], "dir": str(TMP)}, Context())
            spec = ocr.call_args.kwargs["pages"]
            self.assertEqual(len(spec.split(",")), 4001)
            self.assertTrue(spec.endswith(",4001"))

    def test_crash_between_move_and_output_record_is_recovered(self):
        from zhanlu import config, jobs
        from zhanlu.tasks import TaskHub

        record = TMP / "pending-output-tasks.json"
        partial = TMP / "pending-output-staged.txt"
        final = TMP / "pending-output-final.txt"
        delivered_id = []

        def builder(_params, ctx):
            partial.write_text("owned by this task", encoding="utf-8")
            ctx.reserve_output(partial, final)
            partial.replace(final)
            delivered_id.append(jobs.store.add_output(final)["id"])
            raise SystemExit(9)  # 文件已移到目标，尚未来得及 ctx.output

        with patch.object(config, "TASKS_FILE", record), patch("zhanlu.tasks.cleanup_stale_temps"):
            hub = TaskHub(workers=1)
            hub.register("pending-test", lambda _p: "pending-test", builder)
            task = hub.submit("pending-test", {})
            hub._pool.shutdown(wait=True)
            saved = next(row for row in config.load_json(record, []) if row["id"] == task["id"])
            self.assertEqual(saved["outputs"], [])
            self.assertEqual(len(saved["pending_outputs"]), 1)
            self.assertTrue(final.exists())

            recovered = TaskHub(workers=1)
            recovered.cleanup_pending = jobs._cleanup_pending
            recovered.cleanup_outputs = jobs._cleanup_outputs
            recovered.reconcile_interrupted()
            self.assertFalse(final.exists())
            self.assertIsNone(jobs.store.get(delivered_id[0]))
            self.assertEqual(recovered.get(task["id"])["pending_outputs"], [])
            recovered._pool.shutdown(wait=True)

    def test_changed_pending_file_is_not_deleted(self):
        from zhanlu import jobs
        from zhanlu.tasks import TaskContext

        staged = TMP / "changed-pending-staged.txt"
        final = TMP / "changed-pending-final.txt"
        staged.write_text("original", encoding="utf-8")
        task = {"outputs": [], "pending_outputs": [], "cancel": threading.Event()}
        ctx = TaskContext(task)
        ctx.reserve_output(staged, final)
        staged.replace(final)
        final.write_text("someone changed this", encoding="utf-8")
        try:
            self.assertEqual(jobs._cleanup_pending(task["pending_outputs"]), task["pending_outputs"])
            self.assertTrue(final.exists())
        finally:
            final.unlink(missing_ok=True)

    def test_interrupted_published_output_is_rolled_back(self):
        from zhanlu import config
        from zhanlu.tasks import TaskHub

        record = TMP / "interrupted-output-tasks.json"
        partial = TMP / "interrupted-output.txt"

        def builder(_params, ctx):
            partial.write_text("partial", encoding="utf-8")
            ctx.output({"id": "test-output", "name": partial.name, "path": str(partial), "kind": "txt"})
            raise SystemExit(9)  # 模拟进程在交付文件后突然退出

        with patch.object(config, "TASKS_FILE", record), patch("zhanlu.tasks.cleanup_stale_temps"):
            hub = TaskHub(workers=1)
            hub.register("interrupt-test", lambda _p: "interrupt-test", builder)
            task = hub.submit("interrupt-test", {})
            hub._pool.shutdown(wait=True)
            saved = config.load_json(record, [])
            self.assertEqual(saved[0]["outputs"][0]["id"], "test-output")

            recovered = TaskHub(workers=1)
            self.assertEqual(recovered.get(task["id"])["status"], "interrupted")
            recovered.cleanup_outputs = lambda rows: [Path(row["path"]).unlink() for row in rows]
            recovered.reconcile_interrupted()
            self.assertFalse(partial.exists())
            self.assertEqual(recovered.get(task["id"])["outputs"], [])
            self.assertEqual(config.load_json(record, [])[0]["status"], "interrupted")
            recovered._pool.shutdown(wait=True)

    def test_cancel_just_before_builder_returns_is_not_done(self):
        from zhanlu import config
        from zhanlu.tasks import TaskHub

        started, release = threading.Event(), threading.Event()

        def builder(_params, _ctx):
            started.set()
            release.wait(5)
            return {"file": "should-not-be-reported"}

        with patch.object(config, "TASKS_FILE", TMP / "cancel-race-tasks.json"), \
                patch("zhanlu.tasks.cleanup_stale_temps"):
            hub = TaskHub(workers=1)
            hub.register("race", lambda _p: "race", builder)
            task = hub.submit("race", {})
            self.assertTrue(started.wait(5))
            self.assertTrue(hub.cancel(task["id"]))
            release.set()
            for _ in range(100):
                current = hub.get(task["id"])
                if current["status"] == "cancelled":
                    break
                threading.Event().wait(0.01)
            self.assertEqual(current["status"], "cancelled")
            self.assertIsNone(current["result"])
            hub._pool.shutdown(wait=True)


class TestChannelPreference(unittest.TestCase):
    """ai.prefer 配置项：显式 active 最优先，否则按 prefer 选通道、另一通道兜底。"""

    def setUp(self):
        from zhanlu import config
        from zhanlu.core import ai

        self.s = config.load_settings()
        self.s["ai"] = {
            "configs": [
                {"id": "cloud1", "name": "云", "protocol": "openai", "baseurl": "https://api.example.com/v1", "model": "m1"},
                {"id": "local1", "name": "本地", "protocol": "ollama", "baseurl": "http://127.0.0.1:11434/v1", "model": "m2"},
            ],
            "active": None, "prefer": "cloud",
        }
        config.save_settings(self.s)

    def tearDown(self):
        from zhanlu import config
        config.save_settings(self.s)

    def _set(self, **kw):
        from zhanlu import config
        s = config.load_settings()
        s["ai"].update(kw)
        config.save_settings(s)

    def test_prefer_cloud_picks_remote(self):
        from zhanlu.core import ai
        self.assertEqual(ai.active_config()["id"], "cloud1")

    def test_prefer_local_picks_local(self):
        from zhanlu.core import ai
        self._set(prefer="local")
        self.assertEqual(ai.active_config()["id"], "local1")

    def test_fallback_when_preferred_channel_missing(self):
        from zhanlu.core import ai
        self._set(prefer="local", configs=[self.s["ai"]["configs"][0]])  # 只剩云
        self.assertEqual(ai.active_config()["id"], "cloud1")

    def test_explicit_active_beats_prefer(self):
        from zhanlu.core import ai
        self._set(prefer="cloud", active="local1")
        self.assertEqual(ai.active_config()["id"], "local1")

    def test_save_configs_persists_prefer(self):
        from zhanlu.core import ai
        ai.save_configs({"configs": self.s["ai"]["configs"], "active": None, "prefer": "local"})
        self.assertEqual(ai._settings()["ai"]["prefer"], "local")


class TestOcrPackPlugin(unittest.TestCase):
    def test_package_fallback_uses_existing_models_without_download(self):
        from zhanlu.core import ocrpdf

        with tempfile.TemporaryDirectory() as directory, \
             patch.object(ocrpdf, "_pkg_models_dir", return_value=Path(directory)), \
             patch.object(ocrpdf, "RapidOCR") as engine:
            ocrpdf._engines.pop("package", None)
            try:
                ocrpdf._get_engine("package")
                self.assertEqual(engine.call_args.kwargs["params"]["Global.model_root_dir"], directory)
            finally:
                ocrpdf._engines.pop("package", None)

    def test_install_requires_user_confirmation(self):
        from zhanlu import jobs

        with self.assertRaises(pdfops.OpError):
            jobs.hub._builders["ocr_pack_install"]({"tier": "standard"},
                                                   _FakeCtx())
        with self.assertRaises(pdfops.OpError):
            jobs.hub._builders["ocr_pack_uninstall"]({}, _FakeCtx())

    def test_tier_detection_from_files(self):
        """档位检测：伪造标准档文件名 → available 且 active=standard。"""
        from zhanlu import config
        from zhanlu.core import ocrpdf

        d = config.OCR_MODELS_DIR
        d.mkdir(parents=True, exist_ok=True)
        made = []
        for f in ocrpdf._TIER_FILES["standard"]:
            target = d / f
            target.write_bytes(b"x")
            made.append(target)
        try:
            status = ocrpdf.plugin_status()
            self.assertTrue(status["tiers"]["standard"])
            self.assertEqual(status["active"], "standard")
            self.assertTrue(ocrpdf.available())
        finally:
            for t in made:
                t.unlink(missing_ok=True)


class _FakeCtx:
    class _Task(dict):
        cancel = threading.Event()

    _task = _Task()
    _temp_root = Path(tempfile.mkdtemp())

    def progress(self, *a, **k):
        pass


class TestServerApi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from http.server import ThreadingHTTPServer

        from zhanlu import config, server as srv

        config.save_settings(config.DEFAULT_SETTINGS)  # 隔离：清空其他测试写入的 AI 配置
        srv.auth.dev_mode = True  # API 冒烟走开发模式（放宽令牌，保留 Host/Origin 校验）
        cls.httpd = ThreadingHTTPServer(("127.0.0.1", 0), srv.Handler)
        cls.port = cls.httpd.server_address[1]
        cls.base = f"http://127.0.0.1:{cls.port}"
        cls.thread = threading.Thread(target=cls.httpd.serve_forever, daemon=True)
        cls.thread.start()

        cls.pdf_bytes = Path(make_pdf(TMP / "api.pdf", 4, outline=True, text=True)).read_bytes()

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown()

    def _get(self, path: str):
        with urllib.request.urlopen(self.base + path, timeout=10) as r:
            return json.loads(r.read().decode("utf-8"))

    def _post(self, path: str, body: dict):
        req = urllib.request.Request(
            self.base + path, data=json.dumps(body).encode("utf-8"),
            headers={"Content-Type": "application/json"}, method="POST")
        with urllib.request.urlopen(req, timeout=30) as r:
            return json.loads(r.read().decode("utf-8"))

    def _upload(self, filename: str, data: bytes) -> dict:
        boundary = "----zhanlu-test"
        head = (
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
            f"Content-Type: application/pdf\r\n\r\n"
        ).encode("utf-8")
        body = head + data + f"\r\n--{boundary}--\r\n".encode()
        req = urllib.request.Request(
            self.base + "/api/files/upload", data=body, method="POST",
            headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
        with urllib.request.urlopen(req, timeout=30) as r:
            return json.loads(r.read().decode("utf-8"))

    def _wait_task(self, tid: str) -> dict:
        import time

        for _ in range(120):
            t = self._get(f"/api/tasks/{tid}")
            if t["status"] in ("done", "error", "cancelled"):
                return t
            time.sleep(0.1)
        self.fail("任务超时")

    def test_origin_and_host_guards(self):
        """S01：外部 Origin 与非本机 Host 一律 403。"""
        for path in ("/api/app-info",):
            req = urllib.request.Request(self.base + path, headers={"Origin": "http://evil.example.com"})
            with self.assertRaises(urllib.error.HTTPError) as c:
                urllib.request.urlopen(req, timeout=5)
            self.assertEqual(c.exception.code, 403)
            c.exception.close()
            req = urllib.request.Request(self.base + path, headers={"Host": "evil.example.com"})
            with self.assertRaises(urllib.error.HTTPError) as c2:
                urllib.request.urlopen(req, timeout=5)
            self.assertEqual(c2.exception.code, 403)
            c2.exception.close()

    def test_encrypted_pdf_info_keeps_reader_available(self):
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        writer.encrypt("test-password")
        output = io.BytesIO()
        writer.write(output)
        uploaded = self._upload("encrypted-test.pdf", output.getvalue())
        info = self._get(f"/api/files/{uploaded['files'][0]['id']}/info?probe=text")
        self.assertTrue(info["encrypted"])
        self.assertEqual(info["outline"], [])

    def test_strict_token_required(self):
        from zhanlu import server as srv

        old = srv.auth.dev_mode
        srv.auth.dev_mode = False
        try:
            for path in ("/api/files", "/api/tasks", "/api/ai/configs"):
                req = urllib.request.Request(self.base + path)
                with self.assertRaises(urllib.error.HTTPError) as c:
                    urllib.request.urlopen(req, timeout=5)
                self.assertEqual(c.exception.code, 401, path)
                c.exception.close()
            req = urllib.request.Request(
                self.base + "/api/jobs/run", data=b'{}', method="POST",
                headers={"Content-Type": "application/json"},
            )
            with self.assertRaises(urllib.error.HTTPError) as c:
                urllib.request.urlopen(req, timeout=5)
            self.assertEqual(c.exception.code, 401)
            c.exception.close()
            # 令牌放行
            req = urllib.request.Request(self.base + "/api/files",
                                         headers={"X-Zhanlu-Token": srv.auth.token})
            with urllib.request.urlopen(req, timeout=5) as r:
                self.assertEqual(r.status, 200)
        finally:
            srv.auth.dev_mode = old

    def test_unselected_local_path_cannot_be_registered_or_opened(self):
        from zhanlu import server as srv

        with self.assertRaises(urllib.error.HTTPError) as register_error:
            self._post("/api/files/path", {"path": str(TMP / "api.pdf")})
        self.assertEqual(register_error.exception.code, 404)
        register_error.exception.close()

        with patch.object(srv.os, "startfile", create=True) as startfile:
            with self.assertRaises(urllib.error.HTTPError) as open_error:
                self._post("/api/sys/open-path", {"path": str(TMP)})
            self.assertEqual(open_error.exception.code, 400)
            open_error.exception.close()
            startfile.assert_not_called()

    def test_print_no_longer_invokes_external_pdf_app(self):
        from zhanlu import server as srv

        uploaded = self._upload("print-selected.pdf", self.pdf_bytes)
        fid = uploaded["files"][0]["id"]
        with patch.object(srv.os, "startfile", create=True) as startfile:
            with self.assertRaises(urllib.error.HTTPError) as result:
                self._post("/api/sys/print", {"id": fid})
            self.assertEqual(result.exception.code, 404)
            result.exception.close()
            startfile.assert_not_called()

    def test_desktop_open_requests_are_registered_ids_not_http_paths(self):
        from zhanlu import desktop, server as srv
        path = make_pdf(TMP / "桌面打开 空格.pdf", 2)
        desktop.take()
        desktop.receive([str(path)], srv.store)
        result = self._get("/api/desktop/open-files")
        self.assertEqual(result["files"][0]["name"], path.name)
        # path 字段仅供前端悬停区分同名文件；后续请求一律使用 id，不接受 http path
        self.assertEqual(result["files"][0]["path"], str(path))
        self.assertEqual(self._get("/api/desktop/open-files"), {"files": [], "errors": []})
        with self.assertRaises(urllib.error.HTTPError) as error:
            self._post("/api/desktop/open-files", {"path": str(path)})
        self.assertEqual(error.exception.code, 404)
        error.exception.close()

    def test_api_flow(self):
        info = self._get("/api/app-info")
        self.assertIn("display", info)
        self.assertIn("instance", info)

        up = self._upload("api测试.pdf", self.pdf_bytes)
        self.assertEqual(up["files"][0]["kind"], "pdf")
        self.assertEqual(up["files"][0]["pages"], 4)
        fid = up["files"][0]["id"]

        finfo = self._get(f"/api/files/{fid}/info?probe=text")
        self.assertEqual(finfo["pages"], 4)
        self.assertTrue(finfo["outline"])

        with urllib.request.urlopen(self.base + f"/api/files/{fid}/raw", timeout=10) as r:
            self.assertEqual(len(r.read()), len(self.pdf_bytes))

        # 拆分任务：预览 → 运行 → 产物页数正确
        preview = self._post("/api/jobs/preview", {"kind": "split", "params": {"id": fid, "mode": "boundary", "boundary": 2, "stem": "T"}})
        self.assertTrue(preview["ok"], str(preview))
        self.assertEqual([o["pages"] for o in preview["outputs"]], [2, 2])
        partial = self._post("/api/jobs/preview", {"kind": "split", "params": {"id": fid, "mode": "ranges", "ranges": "1-2", "stem": "T-partial"}})
        self.assertTrue(partial["ok"], str(partial))
        self.assertIn("共 2 页", partial["summary"])
        run = self._post("/api/jobs/run", {"kind": "split", "params": {"id": fid, "mode": "boundary", "boundary": 2, "stem": "T"}})
        t = self._wait_task(run["id"])
        self.assertEqual(t["status"], "done", str(t.get("error")))
        self.assertEqual(len(t["outputs"]), 2)
        registered = {f["id"]: f for f in self._get("/api/files")["files"]}
        self.assertEqual([registered[o["id"]]["pages"] for o in t["outputs"]], [2, 2])

        # 同名冲突：任何保存位置都自动编号（v0.2），预览展示编号后的最终交付名
        preview2 = self._post("/api/jobs/preview", {"kind": "split", "params": {"id": fid, "mode": "boundary", "boundary": 2, "stem": "T"}})
        self.assertTrue(preview2["ok"], str(preview2))
        numbered = [o["name"] for o in preview2["outputs"]]
        self.assertEqual(numbered, ["T-01-p001-p002（1）.pdf", "T-02-p003-p004（1）.pdf"])
        run2 = self._post("/api/jobs/run", {"kind": "split", "params": {"id": fid, "mode": "boundary", "boundary": 2, "stem": "T"}})
        t3 = self._wait_task(run2["id"])
        self.assertEqual(t3["status"], "done", str(t3.get("error")))
        self.assertEqual(sorted(Path(o["path"]).name for o in t3["outputs"]),
                         ["T-01-p001-p002（1）.pdf", "T-02-p003-p004（1）.pdf"])

        # PDF→TXT 转换
        conv = self._post("/api/jobs/run", {"kind": "convert_pdf", "params": {"id": fid, "target": "txt", "prefix": "T"}})
        t2 = self._wait_task(conv["id"])
        self.assertEqual(t2["status"], "done", str(t2.get("error")))
        self.assertEqual(len(t2["outputs"]), 1)

        # 任务持久化：重启后可查
        from zhanlu import jobs as jobs_mod
        rows = jobs_mod.hub.list()
        self.assertTrue(any(x["id"] == run["id"] for x in rows))

        engines = self._get("/api/engines")
        self.assertIn("office", engines)

    def _register_local(self, paths) -> list[dict]:
        """走桌面通道登记真实路径，模拟受信任的本机打开。"""
        from zhanlu import desktop, server as srv

        desktop.take()
        desktop.receive([str(p) for p in paths], srv.store)
        return self._get("/api/desktop/open-files")["files"]

    def test_convert_same_source_dir_auto_numbering(self):
        """06 改版 107–115：与源文件同目录 + 同名自动编号，绝不覆盖。"""
        src_dir = TMP / "同目录编号"
        src_dir.mkdir(parents=True, exist_ok=True)
        (fid,) = [f["id"] for f in self._register_local([make_pdf(src_dir / "报告.pdf", 2, text=True)])]

        prev = self._post("/api/jobs/preview", {"kind": "convert_pdf",
                                                "params": {"id": fid, "target": "txt", "prefix": "报告", "dir": "@source"}})
        self.assertTrue(prev["ok"], str(prev))
        self.assertTrue(prev["numbered"])
        self.assertEqual(prev["dir"], str(src_dir.resolve()))
        self.assertEqual(prev["outputs"][0]["name"], "报告.txt")

        run = self._post("/api/jobs/run", {"kind": "convert_pdf",
                                           "params": {"id": fid, "target": "txt", "prefix": "报告", "dir": "@source"}})
        t = self._wait_task(run["id"])
        self.assertEqual(t["status"], "done", str(t.get("error")))
        self.assertTrue((src_dir / "报告.txt").exists())
        first_bytes = (src_dir / "报告.txt").read_bytes()

        # 第二次：预览与执行都递增编号，已有文件保持原样
        prev2 = self._post("/api/jobs/preview", {"kind": "convert_pdf",
                                                 "params": {"id": fid, "target": "txt", "prefix": "报告", "dir": "@source"}})
        self.assertTrue(prev2["ok"], str(prev2))
        self.assertEqual(prev2["outputs"][0]["name"], "报告（1）.txt")
        run2 = self._post("/api/jobs/run", {"kind": "convert_pdf",
                                            "params": {"id": fid, "target": "txt", "prefix": "报告", "dir": "@source"}})
        t2 = self._wait_task(run2["id"])
        self.assertEqual(t2["status"] == "done", True, str(t2.get("error")))
        self.assertTrue((src_dir / "报告（1）.txt").exists())
        self.assertEqual((src_dir / "报告.txt").read_bytes(), first_bytes)

        # 已有（1）占位时跳到（2）
        (src_dir / "报告（1）.txt").unlink()
        (src_dir / "报告（1）.txt").write_bytes("占位".encode("utf-8"))
        prev3 = self._post("/api/jobs/preview", {"kind": "convert_pdf",
                                                 "params": {"id": fid, "target": "txt", "prefix": "报告", "dir": "@source"}})
        self.assertEqual(prev3["outputs"][0]["name"], "报告（2）.txt")

    def test_convert_images_source_dir_numbers_every_output(self):
        src_dir = TMP / "同目录图片"
        src_dir.mkdir(parents=True, exist_ok=True)
        (fid,) = [f["id"] for f in self._register_local([make_pdf(src_dir / "画册.pdf", 2)])]
        params = {"id": fid, "target": "png", "dpi": 96, "prefix": "画册", "pages": "1-2", "dir": "@source"}

        run = self._post("/api/jobs/run", {"kind": "convert_pdf", "params": params})
        t = self._wait_task(run["id"])
        self.assertEqual(t["status"], "done", str(t.get("error")))
        self.assertTrue((src_dir / "画册-p001.png").exists())
        self.assertTrue((src_dir / "画册-p002.png").exists())

        prev = self._post("/api/jobs/preview", {"kind": "convert_pdf", "params": params})
        self.assertTrue(prev["ok"], str(prev))
        self.assertEqual([o["name"] for o in prev["outputs"]], ["画册-p001（1）.png", "画册-p002（1）.png"])
        run2 = self._post("/api/jobs/run", {"kind": "convert_pdf", "params": params})
        t2 = self._wait_task(run2["id"])
        self.assertEqual(t2["status"], "done", str(t2.get("error")))
        self.assertTrue((src_dir / "画册-p001（1）.png").exists())
        self.assertTrue((src_dir / "画册-p002（1）.png").exists())
        self.assertGreater((src_dir / "画册-p001.png").stat().st_size, 0)

    def test_uploaded_inbox_copy_cannot_use_source_dir(self):
        """浏览器上传的 inbox 副本没有源目录（114）：预览说明原因且不通过。"""
        up = self._upload("上传副本.pdf", self.pdf_bytes)
        self.assertIsNone(up["files"][0]["source_dir"])
        prev = self._post("/api/jobs/preview", {"kind": "convert_pdf",
                                                "params": {"id": up["files"][0]["id"], "target": "txt",
                                                           "prefix": "上传副本", "dir": "@source"}})
        self.assertFalse(prev["ok"])
        self.assertIn("与源文件同目录不可用", prev["error"])

    def test_txt_conversion_rejects_selected_pages_without_text(self):
        blank_pdf = make_pdf(TMP / "blank-for-txt.pdf", 2)
        uploaded = self._upload("blank-for-txt.pdf", blank_pdf.read_bytes())
        result = self._post("/api/jobs/run", {"kind": "convert_pdf", "params": {
            "id": uploaded["files"][0]["id"], "target": "txt", "prefix": "blank-for-txt"}})
        task = self._wait_task(result["id"])
        self.assertEqual(task["status"], "error")
        self.assertIn("请先完成 OCR", task["error"])
        self.assertEqual(task["outputs"], [])

    def test_merge_source_dir_follows_first_item(self):
        """合并以当前列表第一项的源目录为准；同名自动编号。"""
        dir_a, dir_b = TMP / "合并甲", TMP / "合并乙"
        dir_a.mkdir(parents=True, exist_ok=True)
        dir_b.mkdir(parents=True, exist_ok=True)
        entries = self._register_local([make_pdf(dir_a / "甲.pdf", 2), make_pdf(dir_b / "乙.pdf", 2)])
        ids = {e["name"]: e["id"] for e in entries}

        prev = self._post("/api/jobs/preview", {"kind": "merge", "params": {
            "items": [{"id": ids["甲.pdf"]}, {"id": ids["乙.pdf"]}], "name": "合集.pdf", "dir": "@source"}})
        self.assertTrue(prev["ok"], str(prev))
        self.assertEqual(prev["dir"], str(dir_a.resolve()))
        self.assertEqual(prev["outputs"][0]["name"], "合集.pdf")

        run = self._post("/api/jobs/run", {"kind": "merge", "params": {
            "items": [{"id": ids["甲.pdf"]}, {"id": ids["乙.pdf"]}], "name": "合集.pdf", "dir": "@source"}})
        t = self._wait_task(run["id"])
        self.assertEqual(t["status"], "done", str(t.get("error")))
        self.assertEqual(len(PdfReader(str(dir_a / "合集.pdf")).pages), 4)
        self.assertFalse((dir_b / "合集.pdf").exists())

    def test_split_source_dir_numbers_outputs(self):
        """2026-10-30：拆分同样可选与源文件同目录；多份输出重名自动编号，不覆盖已有文件。"""
        dir_a = TMP / "拆分同目录"
        dir_a.mkdir(parents=True, exist_ok=True)
        (fid,) = [f["id"] for f in self._register_local([make_pdf(dir_a / "承诺书.pdf", 4)])]

        # 预占第一份产物的同名文件，模拟源目录里已有上次拆分结果
        (dir_a / "承诺书-01-p001-p002.pdf").write_bytes(self.pdf_bytes)

        prev = self._post("/api/jobs/preview", {"kind": "split", "params": {
            "id": fid, "mode": "every", "every": 2, "stem": "承诺书", "dir": "@source"}})
        self.assertTrue(prev["ok"], str(prev))
        self.assertTrue(prev["numbered"], str(prev))
        self.assertEqual(prev["dir"], str(dir_a.resolve()))
        self.assertEqual([o["name"] for o in prev["outputs"]],
                         ["承诺书-01-p001-p002（1）.pdf", "承诺书-02-p003-p004.pdf"])

        run = self._post("/api/jobs/run", {"kind": "split", "params": {
            "id": fid, "mode": "every", "every": 2, "stem": "承诺书", "dir": "@source"}})
        t = self._wait_task(run["id"])
        self.assertEqual(t["status"], "done", str(t.get("error")))
        self.assertTrue((dir_a / "承诺书-01-p001-p002（1）.pdf").exists(), str(t))
        self.assertTrue((dir_a / "承诺书-02-p003-p004.pdf").exists(), str(t))
        self.assertEqual(len(PdfReader(str(dir_a / "承诺书-01-p001-p002（1）.pdf")).pages), 2)
        # 已有文件与源文件都未被改动
        self.assertEqual((dir_a / "承诺书-01-p001-p002.pdf").read_bytes(), self.pdf_bytes)
        self.assertEqual(len(PdfReader(str(dir_a / "承诺书.pdf")).pages), 4)

    def test_register_picked_reuses_entry_and_rejects_non_pdf(self):
        """原生对话框选择的路径按可信通道登记：真实 source_dir、重复选择复用、非 PDF 拒绝。"""
        from zhanlu import server as srv

        path = make_pdf(TMP / "原生选择.pdf", 2)
        entries, errors = srv._register_picked([str(path)])
        self.assertEqual(errors, [])
        self.assertEqual(entries[0]["path"], str(path))
        self.assertEqual(entries[0]["pages"], 2)
        self.assertEqual(str(srv.source_dir_of(entries[0])), str(path.parent))

        again, _ = srv._register_picked([str(path)])
        self.assertEqual(again[0]["id"], entries[0]["id"])  # 同一文件只登记一次

        txt = path.with_suffix(".txt")
        txt.write_text("not a pdf", encoding="utf-8")
        _, errs = srv._register_picked([str(txt)])
        self.assertTrue(errs and "PDF" in errs[0])

    def test_convert_picker_registers_real_image_path(self):
        """转换原生选择支持图片并保留真实目录，不接受任意类型。"""
        from types import SimpleNamespace
        from zhanlu import server as srv

        image = TMP / "原生选择图片.png"
        image.write_bytes(b"image test")
        picked = SimpleNamespace(returncode=0, stdout=str(image) + "\n")
        with patch.object(srv.subprocess, "run", return_value=picked) as run:
            result = srv.api_pick_pdfs({}, {"kind": "convert"})
        self.assertFalse(result["cancelled"])
        self.assertEqual(result["files"][0]["kind"], "image")
        self.assertEqual(result["files"][0]["source_dir"], str(image.parent))
        self.assertIn("*.png", run.call_args.args[0][-1])
        with self.assertRaises(srv.ApiError):
            srv.api_pick_pdfs({}, {"kind": "unknown"})

    def test_picker_reports_dialog_failure(self):
        from types import SimpleNamespace
        from zhanlu import server as srv

        with patch.object(srv.subprocess, "run", return_value=SimpleNamespace(returncode=1, stdout="")):
            with self.assertRaises(srv.ApiError):
                srv.api_pick_pdfs({}, {})

    def test_ai_remote_requires_consent(self):
        from zhanlu import jobs as jobs_mod

        up = self._upload("remote-consent.pdf", self.pdf_bytes)
        cfg = {"id": "rc", "name": "远程测试", "protocol": "openai", "baseurl": "https://example.com/v1",
               "model": "test", "key_plain": "sk-x"}
        from zhanlu.core import ai as ai_mod
        with patch.object(ai_mod, "active_config_full", return_value=cfg):
            run = self._post("/api/jobs/run", {"kind": "ai_analyze", "params": {"id": up["files"][0]["id"], "action": "summary"}})
            t = self._wait_task(run["id"])
        self.assertEqual(t["status"], "error")
        self.assertIn("远程", t["error"])


if __name__ == "__main__":
    unittest.main()


class TestTaskClearFinished(unittest.TestCase):
    """任务弹层「全部清理」：只清终态记录，运行中任务保留。"""

    def test_clear_finished_keeps_running_and_returns_count(self):
        from unittest.mock import patch
        from zhanlu import config
        from zhanlu.tasks import TaskHub

        record = TMP / "clear-finished-tasks.json"

        def builder(params, ctx):
            return {"ok": True}

        with patch.object(config, "TASKS_FILE", record), patch("zhanlu.tasks.cleanup_stale_temps"):
            hub = TaskHub(workers=1)
            hub.register("clear-test", lambda _p: "clear-test", builder)
            done = hub.submit("clear-test", {})
            hub._pool.shutdown(wait=True)
            # 直接注入一条 running 记录（绕过执行器，只测清理语义）
            hub._tasks["fake-running"] = {"id": "fake-running", "kind": "clear-test",
                                          "name": "运行中", "params": {}, "status": "running",
                                          "progress": 0.5, "created": 0, "outputs": [],
                                          "result": None}
            hub._tasks["fake-cancelling"] = {**hub._tasks["fake-running"], "id": "fake-cancelling", "status": "cancelling"}
            hub._tasks["fake-pending"] = {**hub._tasks["fake-running"], "id": "fake-pending", "status": "pending"}
            self.assertFalse(hub.clear("fake-cancelling"))
            count = hub.clear_finished()
            self.assertEqual(count, 1)  # 只有 done 一条被清
            self.assertNotIn(done["id"], hub._tasks)
            self.assertIn("fake-running", hub._tasks)
            self.assertIn("fake-cancelling", hub._tasks)
            self.assertIn("fake-pending", hub._tasks)
            # 再清一次：没有终态记录，返回 0
            self.assertEqual(hub.clear_finished(), 0)
            hub._pool.shutdown(wait=True)


class TestProgressKnownProtocol(unittest.TestCase):
    """剑形进度数据协议：progress(None) 标记不可估算并保留最后已知值；数值恢复后重新可确定；真实 0% 不被误判。"""

    def test_none_marks_unknown_and_value_recovers_known(self):
        from unittest.mock import patch
        from zhanlu import config
        from zhanlu.tasks import TaskHub

        record = TMP / "progress-known-tasks.json"
        seen = {}

        def builder(params, ctx):
            ctx.progress(0.5)
            ctx.progress(None, "等待上游输出")
            seen["unknown"] = {"known": ctx._task["progress_known"], "value": ctx._task["progress"]}
            seen["public_unknown"] = hub._public(ctx._task)["progressKnown"]
            ctx.progress(0.0)  # 真实回到 0 也必须重新视为可确定
            seen["zero"] = {"known": ctx._task["progress_known"], "value": ctx._task["progress"]}
            return {"ok": True}

        with patch.object(config, "TASKS_FILE", record), patch("zhanlu.tasks.cleanup_stale_temps"):
            hub = TaskHub(workers=1)
            hub.register("pk-test", lambda _p: "pk-test", builder)
            task = hub.submit("pk-test", {})
            hub._pool.shutdown(wait=True)
            self.assertEqual(seen["unknown"], {"known": False, "value": 0.5})
            self.assertIs(seen["public_unknown"], False)
            self.assertEqual(seen["zero"], {"known": True, "value": 0.0})
            saved = config.load_json(record, [])
            row = next(r for r in saved if r["id"] == task["id"])
            self.assertIn("progress_known", row)
            self.assertIs(hub.get(task["id"])["progressKnown"], True)



class UploadDedupeTest(unittest.TestCase):
    """同一文件仅能打开一次：上传按内容 sha256 去重（含旧条目回算）；产物登记不受影响。"""

    @classmethod
    def setUpClass(cls):
        from http.server import ThreadingHTTPServer

        from zhanlu import config, server as srv

        config.save_settings(config.DEFAULT_SETTINGS)
        srv.auth.dev_mode = True
        cls.httpd = ThreadingHTTPServer(("127.0.0.1", 0), srv.Handler)
        cls.base = f"http://127.0.0.1:{cls.httpd.server_address[1]}"
        cls.store = srv.store
        threading.Thread(target=cls.httpd.serve_forever, daemon=True).start()
        cls.pdf_bytes = Path(make_pdf(TMP / "dedupe.pdf", 2)).read_bytes()

    @classmethod
    def tearDownClass(cls):
        cls.httpd.shutdown()

    def _upload(self, filename: str, data: bytes) -> dict:
        boundary = "----zhanlu-test"
        head = (
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
            f"Content-Type: application/pdf\r\n\r\n"
        ).encode("utf-8")
        body = head + data + f"\r\n--{boundary}--\r\n".encode()
        req = urllib.request.Request(
            self.base + "/api/files/upload", data=body, method="POST",
            headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
        with urllib.request.urlopen(req, timeout=30) as r:
            return json.loads(r.read().decode("utf-8"))

    def test_same_content_returns_same_entry(self):
        a = self._upload("同文件.pdf", self.pdf_bytes)
        b = self._upload("同文件-再传.pdf", self.pdf_bytes)
        self.assertEqual(a["files"][0]["id"], b["files"][0]["id"])

    def test_different_content_gets_new_entry(self):
        a = self._upload("甲.pdf", self.pdf_bytes)
        c = self._upload("乙.pdf", self.pdf_bytes + b"%tail")
        self.assertNotEqual(a["files"][0]["id"], c["files"][0]["id"])

    def test_dedupe_matches_legacy_entries_without_sha(self):
        # 旧登记（无 sha 字段）按文件回算匹配：已登记文件再上传命中原条目
        from zhanlu import config

        path = make_pdf(config.INBOX_DIR / "legacy-同内容.pdf", 1)
        legacy = self.store.add_path(path)
        self.assertNotIn("sha", legacy)
        up = self._upload("legacy-同内容.pdf", path.read_bytes())
        self.assertEqual(up["files"][0]["id"], legacy["id"])

    def test_output_registration_still_allows_identical_bytes(self):
        # 任务产物不走去重：相同字节的两个产物各自登记
        path = make_pdf(TMP / "output-同内容.pdf", 1)
        one = self.store.add_output(path, name="产物一.pdf")
        two = self.store.add_output(path, name="产物二.pdf")
        self.assertNotEqual(one["id"], two["id"])


class TestAiOutline(unittest.TestCase):
    """AI 识别目录：outline 动作的发送范围与结构化解析（无书签文档的 AI 目录）。"""

    @classmethod
    def setUpClass(cls):
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate
        from zhanlu.core.topdf import cjk_font

        cls.dir = TMP / "ai-outline"
        cls.dir.mkdir(parents=True, exist_ok=True)
        cls.pdf = cls.dir / "无目录.pdf"
        head = ParagraphStyle("head", fontName=cjk_font(), fontSize=16, leading=22)
        body = ParagraphStyle("body", fontName=cjk_font(), fontSize=11, leading=16)
        doc = SimpleDocTemplate(str(cls.pdf), pagesize=A4)
        doc.build([
            Paragraph("一、总体方案", head), Paragraph("背景与目标说明。", body), PageBreak(),
            Paragraph("1.1 建设目标", head), Paragraph("目标细项。", body), PageBreak(),
            Paragraph("二、实施计划", head), Paragraph("计划说明。", body),
        ])

    def test_scope_covers_all_text_pages(self):
        from zhanlu.core import document_ai

        scope = document_ai.sending_scope(self.pdf, "outline")
        self.assertEqual(scope["pages"], [1, 2, 3])
        self.assertGreater(scope["characters"], 0)
        self.assertIn("目录", scope["detail"])

    def test_parse_drops_invalid_and_dedupes(self):
        from zhanlu.core.document_ai import _outline_from_answer

        raw = ('前置说明\n```json\n[{"t":"一、总体方案","p":1,"l":1},{"t":"1.1 建设目标","p":2,"l":2},'
               '{"t":"页码越界","p":99},{"t":"","p":2},{"t":"1.1 建设目标","p":2,"l":2},'
               '{"t":"层级噪声","p":3,"l":5}]\n```')
        self.assertEqual(_outline_from_answer(raw, 3), [
            {"title": "一、总体方案", "page": 1, "level": 1},
            {"title": "1.1 建设目标", "page": 2, "level": 2},
            {"title": "层级噪声", "page": 3, "level": 1},
        ])

    def test_analyze_outline_returns_structured_entries(self):
        from unittest.mock import patch

        from zhanlu.core import ai as ai_mod, document_ai

        cfg = {"id": "l", "name": "本地", "protocol": "ollama",
               "baseurl": "http://127.0.0.1:11434/v1", "model": "m"}
        reply = '[{"t":"一、总体方案","p":1,"l":1},{"t":"1.1 建设目标","p":2,"l":2}]'
        with patch.object(ai_mod, "chat", return_value=reply) as mocked:
            result = document_ai.analyze(self.pdf, cfg, "outline")
        self.assertEqual(result["outline"], [
            {"title": "一、总体方案", "page": 1, "level": 1},
            {"title": "1.1 建设目标", "page": 2, "level": 2}])
        self.assertTrue(result["local"])
        # 摘录必须带页码锚点，模型才能给出真实页码
        self.assertIn("[第 1 页]", mocked.call_args[0][1][1]["content"])

    def test_analyze_outline_unparseable_raises(self):
        from unittest.mock import patch

        from zhanlu.core import ai as ai_mod, document_ai

        cfg = {"id": "l", "protocol": "ollama", "baseurl": "http://127.0.0.1:11434/v1", "model": "m"}
        with patch.object(ai_mod, "chat", return_value="抱歉，我无法输出目录。"):
            with self.assertRaises(ValueError):
                document_ai.analyze(self.pdf, cfg, "outline")
