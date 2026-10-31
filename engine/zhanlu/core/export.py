"""导出为 Office：PDF → DOCX（文本流基础版）、PDF → PPTX（页面图片模式）。"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Callable

try:
    import docx  # noqa: F401
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import pptx  # noqa: F401
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

from .pdfops import OpError, extract_text, outline, parse_pages, _reader

# PDF 提取文本常混入 XML 不接受的控制字符（垂直 tab、设备控制符等），写入前清理
_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _clean(text: str) -> str:
    return _XML_ILLEGAL.sub("", text)


def pdf_to_docx(path: Path, out: Path, password: str | None = None,
                pages: str | None = None,
                progress: Callable[[int, int], None] | None = None) -> Path:
    """文本流基础版：每页文字按行成段，页间分页符；含书签时生成目录。"""
    try:
        from docx import Document
        from docx.enum.text import WD_BREAK
    except ImportError as exc:
        raise OpError("DOCX 导出未安装（python-docx 缺失）") from exc

    reader = _reader(path, password)
    meta = reader.metadata or {}
    title = str(meta.get("/Title") or path.stem)
    nodes = outline(path, password)
    pages_text = extract_text(path, pages=pages, password=password)
    total = len(pages_text)

    doc = Document()
    h = doc.add_heading(_clean(title), level=0)
    if nodes:
        doc.add_heading("目录", level=1)
        _write_outline(doc, nodes)
        doc.add_paragraph()

    for pi, text in enumerate(pages_text, 1):
        blocks = [b.strip() for b in (text or "").splitlines() if b.strip()]
        for b in blocks:
            doc.add_paragraph(_clean(b))
        if pi != total:
            if doc.paragraphs:
                doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        if progress is not None:
            progress(pi, total)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def _write_outline(doc, nodes, depth=0):
    from docx.shared import Pt

    for n in nodes:
        page = f"　· {n['page']} 页" if n.get("page") else ""
        p = doc.add_paragraph()
        run = p.add_run(_clean(("　" * depth) + ("• " if depth else "") + n["title"] + page))
        run.font.size = Pt(12 if depth else 13)
        if n.get("children"):
            _write_outline(doc, n["children"], depth + 1)


def pdf_to_pptx(image_paths: list[Path], out: Path) -> Path:
    """每个页面图片一页（16:9 全幅）——图片模式转 PPT。"""
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError as exc:
        raise OpError("PPTX 导出未安装（python-pptx 缺失）") from exc
    if not image_paths:
        raise OpError("没有页面图片，无法生成 PPT")

    prs = Presentation()
    prs.slide_width = Emu(12192000)  # 16:9
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]
    for img in image_paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
